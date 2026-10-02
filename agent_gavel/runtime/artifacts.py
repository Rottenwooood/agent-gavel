"""产物存储（docs/tech-plan.md §5.9）。

下载/截图/trace 等产物统一存到 artifact 目录，用 id 引用；模型拿不到任意
绝对路径，要落盘必须走 export（受文件白名单约束）。
M1 提供基础存储与清理；下载捕获、MIME/文本抽取在 M3 补全。
"""

import hashlib
import json
import os
import shutil
import tempfile
import time

from .errors import GavelError

_DATA_ROOT = os.environ.get(
    "AGENT_GAVEL_DATA_DIR",
    os.path.join(tempfile.gettempdir(), "agent-gavel"),
)
ARTIFACT_DIR = os.environ.get(
    "AGENT_GAVEL_ARTIFACT_DIR", os.path.join(_DATA_ROOT, "artifacts"))


def _now():
    return time.time()


class ArtifactStore:
    def __init__(self, root: str = ARTIFACT_DIR):
        self.root = root
        self._index_path = os.path.join(root, "index.json")
        self._index = None

    # ---- index ----
    def _load(self):
        if self._index is not None:
            return self._index
        os.makedirs(self.root, exist_ok=True)
        try:
            with open(self._index_path, encoding="utf-8") as f:
                self._index = json.load(f)
        except Exception:
            self._index = {}
        return self._index

    def _save(self):
        os.makedirs(self.root, exist_ok=True)
        tmp = self._index_path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self._index or {}, f, ensure_ascii=False, indent=2)
        os.replace(tmp, self._index_path)

    def _next_id(self):
        idx = self._load()
        n = len(idx) + 1
        while f"artifact_{n:03d}" in idx:
            n += 1
        return f"artifact_{n:03d}"

    # ---- 写入 ----
    def save_bytes(self, data: bytes, *, filename=None, kind="file",
                   mime_type=None, source_url=None, session_id=None,
                   page_id=None, meta=None) -> dict:
        aid = self._next_id()
        dest = os.path.join(self.root, aid)
        with open(dest, "wb") as f:
            f.write(data)
        info = {
            "artifact_id": aid,
            "kind": kind,
            "suggested_filename": filename or aid,
            "mime_type": mime_type,
            "size_bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
            "source_url": source_url,
            "session_id": session_id,
            "page_id": page_id,
            "path": dest,
            "created_at": _now(),
        }
        if meta:
            info.update(meta)
        self._load()[aid] = info
        self._save()
        return dict(info)

    def register_file(self, src_path: str, *, filename=None, kind="file",
                      move=True, **kw) -> dict:
        if not os.path.isfile(src_path):
            raise GavelError("artifact_missing", f"文件不存在：{src_path}")
        aid = self._next_id()
        dest = os.path.join(self.root, aid)
        if move:
            shutil.move(src_path, dest)
        else:
            shutil.copy2(src_path, dest)
        size = os.path.getsize(dest)
        h = hashlib.sha256()
        with open(dest, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
        info = {
            "artifact_id": aid,
            "kind": kind,
            "suggested_filename": filename or os.path.basename(src_path) or aid,
            "mime_type": kw.pop("mime_type", None),
            "size_bytes": size,
            "sha256": h.hexdigest(),
            "path": dest,
            "created_at": _now(),
        }
        info.update(kw)
        self._load()[aid] = info
        self._save()
        return dict(info)

    # ---- 读 ----
    def get(self, artifact_id: str) -> dict:
        info = self._load().get(artifact_id)
        if not info:
            raise GavelError("artifact_not_found", f"未知 artifact：{artifact_id}")
        return dict(info)

    def list(self, *, kind=None, run_id=None, session_id=None) -> list:
        out = []
        for info in self._load().values():
            if kind and info.get("kind") != kind:
                continue
            if run_id and info.get("run_id") != run_id:
                continue
            if session_id and info.get("session_id") != session_id:
                continue
            out.append(dict(info))
        return sorted(out, key=lambda x: x.get("created_at", 0))

    @staticmethod
    def _safe_filename(name):
        """只保留纯文件名：剥掉目录分隔 / 盘符 / 空字节，拒绝 . 与 ..。

        防路径穿越：filename=../escaped.txt 会退化成 escaped.txt（写在
        dest_dir 内），filename 为绝对路径同理只取末段。
        """
        raw = str(name or "").replace("\\", "/").replace("\x00", "")
        base = raw.rsplit("/", 1)[-1].strip()
        if os.path.splitdrive(base)[0]:
            raise GavelError("bad_filename", f"非法文件名（含盘符）：{name}")
        if not base or base in (".", ".."):
            raise GavelError("bad_filename", f"非法文件名：{name!r}")
        return base

    def export(self, artifact_id: str, dest_dir: str, filename=None,
               *, policies=None, path_visibility: str = "hidden") -> dict:
        info = self.get(artifact_id)
        if not dest_dir:
            raise GavelError("bad_args", "需要 dest_dir")
        if policies is not None:
            policies.check_local_path(dest_dir, purpose="导出")
        os.makedirs(dest_dir, exist_ok=True)
        name = self._safe_filename(
            filename or info.get("suggested_filename") or artifact_id)
        dest = os.path.join(dest_dir, name)
        # 二次包含校验：净化后的最终路径必须仍落在 dest_dir 内
        root = os.path.realpath(dest_dir)
        parent = os.path.realpath(os.path.dirname(dest))
        if not (parent == root or parent.startswith(root + os.sep)):
            raise GavelError("policy_blocked", f"导出路径越出目标目录：{name}",
                             detail={"dest_dir": dest_dir})
        base, ext = os.path.splitext(dest)
        i = 1
        while os.path.exists(dest):
            dest = f"{base}_{i}{ext}"
            i += 1
        shutil.copy2(info["path"], dest)
        # 默认不回传绝对路径（对齐 artifact_get 的隐藏原则）
        out = {"artifact_id": artifact_id, "filename": os.path.basename(dest),
               "exported": True, "size_bytes": info.get("size_bytes")}
        if path_visibility == "user_visible":
            out["exported_to"] = dest
        return out

    def delete(self, artifact_id: str) -> dict:
        idx = self._load()
        if artifact_id not in idx:
            raise GavelError("artifact_not_found", f"未知 artifact：{artifact_id}")
        self._remove(artifact_id, idx)
        self._save()
        return {"artifact_id": artifact_id, "deleted": True}

    # ---- 下载捕获（page.on("download") 调用）----
    async def capture_download(self, download, *, session_id=None, page_id=None,
                               policies=None, events=None):
        aid = None
        try:
            failure = await download.failure()
            if failure:
                if events:
                    events.emit({"event": "download_failed", "session_id": session_id,
                                 "page_id": page_id, "error": failure,
                                 "suggested_filename": download.suggested_filename})
                return None
            src = await download.path()
            size = os.path.getsize(src)
            if policies is not None:
                try:
                    policies.check_download_size(size)
                except GavelError as e:
                    if events:
                        events.emit({"event": "download_rejected",
                                     "session_id": session_id, "page_id": page_id,
                                     "error": e.message, "size_bytes": size})
                    return None
            aid = self._next_id()
            dest = os.path.join(self.root, aid)
            await download.save_as(dest)
            filename = download.suggested_filename or aid
            info = {
                "artifact_id": aid,
                "kind": "download",
                "suggested_filename": filename,
                "mime_type": _guess_mime(dest, filename),
                "size_bytes": size,
                "sha256": _sha256(dest),
                "source_url": download.url,
                "session_id": session_id,
                "page_id": page_id,
                "path": dest,
                "created_at": _now(),
            }
            self._load()[aid] = info
            self._save()
            if events:
                events.emit({"event": "download_completed", "session_id": session_id,
                             "page_id": page_id, "artifact_id": aid,
                             "suggested_filename": filename,
                             "size_bytes": size, "sha256": info["sha256"]})
            return dict(info)
        except Exception as e:  # noqa: BLE001
            if events:
                events.emit({"event": "download_failed", "session_id": session_id,
                             "page_id": page_id, "error": str(e)})
            return None

    # ---- 文本抽取 ----
    def read_text(self, artifact_id: str) -> dict:
        info = self.get(artifact_id)
        with open(info["path"], "rb") as f:
            data = f.read()
        text, kind = extract_text(data, info.get("mime_type"))
        if text is None:
            return {"artifact_id": artifact_id, "extracted": False,
                    "kind": kind, "hint": "该类型不支持文本抽取"}
        return {"artifact_id": artifact_id, "extracted": True, "kind": kind,
                "text": text, "chars": len(text)}

    # ---- 清理 ----
    def cleanup(self, *, ttl_days=None, max_bytes=None) -> dict:
        idx = self._load()
        removed = []
        now = _now()
        if ttl_days is not None:
            cutoff = now - ttl_days * 86400
            for aid, info in list(idx.items()):
                if info.get("created_at", now) < cutoff:
                    self._remove(aid, idx)
                    removed.append(aid)
        if max_bytes is not None:
            items = sorted(idx.items(), key=lambda kv: kv[1].get("created_at", 0))
            total = sum(i.get("size_bytes", 0) for _, i in items)
            for aid, info in items:
                if total <= max_bytes:
                    break
                total -= info.get("size_bytes", 0)
                self._remove(aid, idx)
                removed.append(aid)
        if removed:
            self._save()
        return {"removed": removed, "count": len(removed)}

    def _remove(self, aid, idx):
        info = idx.pop(aid, None)
        if info and info.get("path") and os.path.isfile(info["path"]):
            try:
                os.remove(info["path"])
            except OSError:
                pass


# ---------------- 模块级工具：哈希 / MIME / 文本抽取 ----------------

def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _guess_mime(path, filename=None):
    name = (filename or path or "").lower()
    ext = os.path.splitext(name)[1]
    table = {".pdf": "application/pdf",
             ".docx": "application/vnd.openxmlformats-officedocument."
                      "wordprocessingml.document",
             ".doc": "application/msword",
             ".xlsx": "application/vnd.openxmlformats-officedocument."
                      "spreadsheetml.sheet",
             ".xls": "application/vnd.ms-excel",
             ".csv": "text/csv", ".txt": "text/plain",
             ".json": "application/json", ".xml": "application/xml",
             ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".gif": "image/gif", ".zip": "application/zip"}
    if ext in table:
        return table[ext]
    try:
        with open(path, "rb") as f:
            head = f.read(8)
        if head[:4] == b"%PDF":
            return "application/pdf"
        if head[:2] == b"PK":
            return "application/zip"
        if head[:4] == b"\xd0\xcf\x11\xe0":
            return "application/msword"
    except Exception:
        pass
    return "application/octet-stream"


def extract_text(data: bytes, content_type=None):
    """按类型抽取文本，返回 (text|None, kind)。"""
    import io
    ct = (content_type or "").lower()
    if data[:4] == b"%PDF" or "pdf" in ct:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages), "pdf"
        except Exception as e:  # noqa: BLE001
            return None, f"pdf-error:{str(e)[:120]}"
    if data[:2] == b"PK" or "wordprocessingml" in ct or "officedocument" in ct:
        try:
            from docx import Document
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append("\t".join(c.text for c in row.cells))
            return "\n".join(parts), "docx"
        except Exception as e:  # noqa: BLE001
            return None, f"docx-error:{str(e)[:120]}"
    if data[:4] == b"\xd0\xcf\x11\xe0" or "msword" in ct:
        return _extract_legacy_doc(data)
    if ct.startswith("text/") or data[:1] in (b"{", b"[", b"<") or b"\n" in data[:200]:
        try:
            return data.decode("utf-8", errors="replace"), "text"
        except Exception:
            pass
    return None, ct or "unknown"


def _extract_legacy_doc(data):
    """老 .doc（OLE）：soffice → antiword → catdoc。"""
    import subprocess
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as f:
        f.write(data)
        path = f.name
    try:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            outdir = tempfile.mkdtemp()
            profile = tempfile.mkdtemp()
            try:
                subprocess.run([soffice, "--headless",
                                f"-env:UserInstallation=file://{profile}",
                                "--convert-to", "txt:Text", "--outdir", outdir, path],
                               capture_output=True, timeout=60)
                out = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".txt")
                if os.path.isfile(out):
                    with open(out, encoding="utf-8", errors="replace") as fh:
                        return fh.read(), "doc"
            except Exception:
                pass
        for tool in ("antiword", "catdoc"):
            exe = shutil.which(tool)
            if not exe:
                continue
            try:
                r = subprocess.run([exe, path], capture_output=True, timeout=30)
                if r.returncode == 0 and r.stdout:
                    return r.stdout.decode("utf-8", errors="replace"), "doc"
            except Exception:
                pass
        return None, "doc-unsupported"
    finally:
        try:
            os.remove(path)
        except OSError:
            pass
