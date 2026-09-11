"""DOM 六项增强端到端测试（MCP 会话）。

覆盖: dom_navigate 元信息/错误页判定、dom_read 三态、dom_text links、
dom_explore 0命中/url、dom_document 文档抽取。

用法: cd agent-gavel && uv run python3 tests/dom_enhancements.py
"""

import asyncio
import functools
import http.server
import json
import os
import sys
import threading
import time

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

HERE = os.path.dirname(os.path.abspath(__file__))
FIXDIR = os.path.join(HERE, "fixtures")
FIX = f"file://{os.path.join(FIXDIR, 'read_text_nav.html')}"
ROOT = os.path.dirname(HERE)
PORT = 8791
BASE = f"http://127.0.0.1:{PORT}"


def make_docx():
    from docx import Document
    p = os.path.join(FIXDIR, "test_doc.docx")
    d = Document()
    d.add_paragraph("河南大学保研政策测试文档 ABC123")
    d.add_paragraph("第二段：推免实施细则。")
    d.save(p)
    return p


def start_http():
    class Handler(http.server.SimpleHTTPRequestHandler):
        def do_GET(self):
            if self.path.startswith("/redirect"):
                self.send_response(302)
                self.send_header("Location", "/read_text_nav.html")
                self.end_headers()
                return
            super().do_GET()

    handler = functools.partial(Handler, directory=FIXDIR)
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", PORT), handler)
    t = threading.Thread(target=httpd.serve_forever, daemon=True)
    t.start()
    return httpd


def parse(res):
    for c in res.content:
        if getattr(c, "type", "") == "text":
            try:
                return json.loads(c.text)
            except Exception:
                return c.text
    return None


async def call(s, name, args, timeout=40):
    t0 = time.perf_counter()
    r = await s.call_tool(name, args)
    return parse(r), int((time.perf_counter() - t0) * 1000)


async def main():
    make_docx()
    httpd = start_http()
    ok_all = True
    params = StdioServerParameters(command=os.path.join(ROOT, "run-mcp.sh"), args=[])
    params.env = {"PATH": os.environ.get("PATH", ""), "HOME": os.environ.get("HOME", ""),
                  "DISPLAY": os.environ.get("DISPLAY", ":0")}
    try:
        async with stdio_client(params) as (rd, wr):
            async with ClientSession(rd, wr) as s:
                await s.initialize()
                tools = [t.name for t in (await s.list_tools()).tools]
                print("新工具:", [t for t in tools if t in ("dom_document",)])

                def check(name, cond, info):
                    nonlocal ok_all
                    if not cond:
                        ok_all = False
                    print(f"[{'OK ' if cond else 'BAD'}] {name:<38} {info}")

                # 1 navigate 元信息（用 http 才能拿 content_type）
                r, ms = await call(s, "dom_navigate", {"url": BASE + "/read_text_nav.html"})
                check("1 navigate 元信息", r.get("navigated") is True
                      and "read_text_nav" in (r.get("final_url") or "")
                      and r.get("content_type") == "text/html"
                      and r.get("status") == "pass",
                      f"{ms}ms navigated={r.get('navigated')} ct={r.get('content_type')}")

                # 2 navigate 404 -> warning/error_page
                r, ms = await call(s, "dom_navigate", {"url": BASE + "/nope_404.html"})
                check("2 navigate 404 -> warning",
                      r.get("status") in ("warning", "fail")
                      and r.get("reason") in ("error_page", "not_navigated"),
                      f"{ms}ms status={r.get('status')} reason={r.get('reason')}")

                # 3 dom_read no_match
                r, ms = await call(s, "dom_read",
                                   {"page_features": {"x": "1"}, "selector": "#nope"})
                check("3 dom_read no_match", r.get("status") == "no_match",
                      f"{ms}ms status={r.get('status')}")

                # 4 dom_read empty（全部为空 -> status=empty）
                r, ms = await call(s, "dom_read", {"page_features": {
                    "e": "document.querySelector('#nope') ? 1 : null"}})
                check("4 dom_read empty", r.get("status") == "empty"
                      and "e" in (r.get("empty") or []),
                      f"{ms}ms status={r.get('status')} empty={r.get('empty')}")

                # 5 dom_read error
                r, ms = await call(s, "dom_read",
                                   {"page_features": {"x": "throw new Error('boom')"}})
                check("5 dom_read error", r.get("status") == "error"
                      and "x" in (r.get("errors") or {}),
                      f"{ms}ms status={r.get('status')} errors={r.get('errors')}")

                # 6 dom_text links（先回到有链接的页面）
                await call(s, "dom_navigate", {"url": FIX})
                r, ms = await call(s, "dom_text", {"mode": "links"})
                links = r.get("links", [])
                has_nav = any("page2.html" in (l.get("href") or "") for l in links)
                check("6 dom_text links", r.get("status") == "ok" and has_nav,
                      f"{ms}ms count={r.get('count')}")

                # 7 dom_explore 0命中 + url 参数
                r, ms = await call(s, "dom_explore", {"url": FIX, "tag": "input"})
                check("7 dom_explore 0命中 hint",
                      r.get("after_filter") == 0 and bool(r.get("hint")),
                      f"{ms}ms after_filter={r.get('after_filter')} hint={r.get('hint')}")

                # 8 dom_document docx 抽取
                r, ms = await call(s, "dom_document",
                                   {"url": BASE + "/test_doc.docx"})
                txt = r.get("text", "") if isinstance(r, dict) else ""
                check("8 dom_document docx", r.get("extracted") is True
                      and "保研政策测试文档" in txt and r.get("kind") == "docx",
                      f"{ms}ms kind={r.get('kind')} len={r.get('length')}")

                # 9 dom_step 拒绝 navigate（导航唯一入口是 dom_navigate）
                r, ms = await call(s, "dom_step",
                                   {"action": "navigate", "selectors": {"url": FIX}})
                check("9 dom_step 拒绝 navigate", r.get("status") == "error"
                      and r.get("reason") == "use_dom_navigate",
                      f"{ms}ms reason={r.get('reason')}")

                # 10 click target=_blank -> 新标签检测/切换
                await call(s, "dom_navigate", {"url": FIX})
                r, ms = await call(s, "dom_step", {
                    "action": "click", "selectors": {"target": "#blank"},
                    "wait_navigation": True, "strict": True})
                nt = r.get("new_tab") or {}
                url = r.get("verification", {}).get("url") or ""
                check("10 click 新标签检测", bool(nt) or "page2" in url,
                      f"{ms}ms new_tab={nt} url={url}")

                # 11 dom_resolve 跳转链（本地 302）
                r, ms = await call(s, "dom_resolve", {"url": BASE + "/redirect"})
                check("11 dom_resolve 跳转链", r.get("status") == "ok"
                      and "read_text_nav" in (r.get("final_url") or ""),
                      f"{ms}ms final={r.get('final_url')}")

                # 12 dom_navigate 断言轮询（延迟 2.5s 出现的特征不应误判 fail）
                r, ms = await call(s, "dom_navigate", {
                    "url": FIX,
                    "page_features": {"d": "document.querySelector('#delayed').textContent"},
                    "expected_feature": {"d": {"op": "eq", "value": "ready"}},
                    "wait_s": 6})
                check("12 navigate 断言轮询", r.get("status") == "pass",
                      f"{ms}ms status={r.get('status')} d={r.get('features',{}).get('d')}")

                # 13 dom_navigate 同主域重定向 -> pass（区域重定向/跟踪参数不误报）
                r, ms = await call(s, "dom_navigate", {"url": BASE + "/redirect"})
                check("13 navigate 同主域重定向", r.get("status") == "pass"
                      and r.get("reached_target") is True,
                      f"{ms}ms status={r.get('status')} reached={r.get('reached_target')}")
    finally:
        httpd.shutdown()

    print("\n" + ("ALL PASS" if ok_all else "HAS FAIL"))
    sys.exit(0 if ok_all else 1)


if __name__ == "__main__":
    asyncio.run(main())
